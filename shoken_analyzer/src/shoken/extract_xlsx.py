"""市販の商圏レポートExcel（1店舗=1ブック）から、ツール用の1行を抜き出す。

ベンダーのブックは店舗ごとに別ファイルなので、ここで横持ちの1行に変換して
既存店マスタに積み上げる。セル番地は config/extract_shoken_report.yaml 側にあるので、
レポートのバージョンが変わってもコードは触らずに済む。
"""

from __future__ import annotations

import re
import unicodedata
import warnings
from pathlib import Path

import openpyxl
import pandas as pd
import yaml

from .config import eval_formula


class ExtractError(ValueError):
    pass


# ベンダーのブックはシート名の末尾に調査年次を付ける（年収特性23W / 世帯特性20W）。
# 年次は店舗ごとに違うことがあるので、照合では落とす。
_VINTAGE = re.compile(r"[0-9]+[A-Za-z]?[0-9]*[A-Za-z]?$")


def _norm_sheet(name) -> str:
    return (unicodedata.normalize("NFKC", str(name))
            .replace(" ", "").replace("\u3000", "").lower())


def _sheet_stem(name) -> str:
    return _VINTAGE.sub("", _norm_sheet(name))


def resolve_sheet(wb, want: str) -> tuple[str | None, str | None]:
    """設定のシート名を、このブックの実際のシート名に対応づける。

    戻り値は (実際のシート名 or None, 利用者に見せる注記 or None)。
    見つからなくても例外にしない。1シート欠けただけで店舗ごと落とすと、
    残り50列以上の正しいデータまで捨てることになるため。
    """
    if want in wb.sheetnames:
        return want, None
    nw = _norm_sheet(want)
    for sh in wb.sheetnames:
        if _norm_sheet(sh) == nw:
            return sh, None
    stem = _sheet_stem(want)
    if stem:
        cands = [sh for sh in wb.sheetnames if _sheet_stem(sh) == stem]
        if len(cands) == 1:
            return cands[0], f"シート '{want}' は '{cands[0]}' として読みました（調査年次違い）"
        if len(cands) > 1:
            pick = max(cands, key=_norm_sheet)      # 年次が新しいものを採る
            return pick, (f"シート '{want}' の候補が複数あります {cands}。"
                          f"'{pick}' を使いました")
    return None, f"シート '{want}' がありません。この項目は空にしました"


def _cell(wb, sheet: str | None, addr: str):
    if sheet is None or sheet not in wb.sheetnames:
        return None
    return wb[sheet][addr].value


def _num(v) -> float | None:
    if v is None or isinstance(v, str):
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return f if pd.notna(f) else None


def _find_by_label(wb, sheet: str | None, label: str, col: str) -> float | None:
    """A列のラベルを前方一致で探し、指定列の値を返す（全角空白は無視）。"""
    if sheet is None or sheet not in wb.sheetnames:
        return None
    ws = wb[sheet]
    norm = label.replace("　", "").replace(" ", "")
    for row in range(1, min(ws.max_row, 200) + 1):
        a = ws[f"A{row}"].value
        if a is None:
            continue
        if str(a).replace("　", "").replace(" ", "").startswith(norm):
            return _num(ws[f"{col}{row}"].value)
    return None


def _find_by_header(wb, sheet: str | None, header: str, header_row: int,
                    data_row: int) -> float | None:
    """ヘッダ行の見出し文字列で列を探し、指定行の値を返す（列ズレに強い）。"""
    if sheet is None or sheet not in wb.sheetnames:
        return None
    ws = wb[sheet]
    norm = header.replace("　", "").replace(" ", "")
    for col in range(1, min(ws.max_column, 400) + 1):
        h = ws.cell(row=header_row, column=col).value
        if h is None:
            continue
        if str(h).replace("　", "").replace(" ", "") == norm:
            return _num(ws.cell(row=data_row, column=col).value)
    return None


from .search_profile import store_id_from_filename, store_name_from_filename

def extract_store(
    xlsx_path: str | Path,
    map_path: str | Path,
    *,
    store_id: str | None = None,
    store_name: str | None = None,
    notes: list[str] | None = None,
) -> dict:
    """1ブック = 1店舗の商圏変数を dict で返す。

    notes を渡すと、シート名の読み替えや欠落をそこに追記する。
    シートが足りなくても例外にはせず、その項目だけ空にして続ける。
    """
    xlsx_path, map_path = Path(xlsx_path), Path(map_path)
    m = yaml.safe_load(map_path.read_text(encoding="utf-8"))

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=False)

    out: dict = {}
    _notes = notes if notes is not None else []
    _resolved: dict[str, str | None] = {}

    def sheet_of(name: str) -> str | None:
        if name not in _resolved:
            actual, note = resolve_sheet(wb, name)
            _resolved[name] = actual
            if note and note not in _notes:
                _notes.append(note)
        return _resolved[name]

    # 店名・商圏定義
    meta = m.get("meta", {})
    name_spec = meta.get("store_name")
    raw_name = _cell(wb, sheet_of(name_spec["sheet"]), name_spec["cell"]) if name_spec else None
    area_spec = meta.get("area_def")
    raw_area = _cell(wb, sheet_of(area_spec["sheet"]), area_spec["cell"]) if area_spec else None
    # ブック内に店名が無ければファイル名から作る。先頭の店舗CDは店名に残さない。
    out["store_name"] = store_name or (str(raw_name).strip() if raw_name
                                       else store_name_from_filename(xlsx_path))
    # ファイル名の頭に店舗CDが付いていればそれを使う（0785_〇〇店.xlsx → 0785）。
    # 店舗名を store_id にすると、IDPOSの店舗CDと突き合わず数字が1つも出なくなる。
    out["store_id"] = store_id or store_id_from_filename(xlsx_path) or out["store_name"]
    out["trade_area_def"] = str(raw_area).strip() if raw_area else ""

    # セル直指定
    for key, spec in (m.get("cells") or {}).items():
        out[key] = _num(_cell(wb, sheet_of(spec["sheet"]), spec["cell"]))

    # ラベル検索（A列の見出しで行を探す）
    for key, spec in (m.get("anchors") or {}).items():
        out[key] = _find_by_label(wb, sheet_of(spec["sheet"]), spec["label"], spec["col"])

    # ヘッダ検索（1行目の見出しで列を探す）
    for key, spec in (m.get("header_lookups") or {}).items():
        out[key] = _find_by_header(
            wb, sheet_of(spec["sheet"]), spec["header"],
            int(spec.get("header_row", 1)), int(spec.get("data_row", 2)),
        )

    # 年齢5歳階級 -> 5区分
    at = m.get("age_table")
    age_sheet = sheet_of(at["sheet"]) if at else None
    if at and age_sheet is None:
        # シートが無くても列は用意しておく（マスタの列を揃えるため）。
        for bucket in at["buckets"]:
            out[bucket] = None
        out["_age_total"] = None
    elif at:
        ws = wb[age_sheet]
        raw: dict[str, float] = {}
        for r in range(int(at["first_row"]), int(at["last_row"]) + 1):
            lbl = ws[f'{at["label_col"]}{r}'].value
            val = _num(ws[f'{at["value_col"]}{r}'].value)
            if lbl is not None and val is not None:
                raw[str(lbl).strip()] = val
        total = sum(raw.values())
        for bucket, labels in at["buckets"].items():
            s = sum(raw.get(l, 0.0) for l in labels)
            out[bucket] = (s / total) if total > 0 else None
        out["_age_total"] = total

    # 計算項目
    for key, formula in (m.get("computed") or {}).items():
        base = {k: v for k, v in out.items() if isinstance(v, (int, float))}
        out[key] = eval_formula(formula, base)

    # このブックからは取れない項目は空で用意しておく（テンプレの列を揃えるため）
    for key in m.get("not_available_in_workbook") or []:
        out.setdefault(key, None)

    wb.close()
    return out


def extract_category_demand(xlsx_path: str | Path, map_path: str | Path) -> pd.DataFrame:
    """商圏のJICFS中分類別 消費推計（1人当たり金額・商圏指数）を取り出す。

    これは「商圏から見たカテゴリ需要の強弱」であって、自店の売上構成比ではない。
    類似度の説明変数には使わず、棚割提案の裏付け・反証に使う。
    """
    m = yaml.safe_load(Path(map_path).read_text(encoding="utf-8"))
    spec = m.get("category_demand")
    if not spec:
        return pd.DataFrame()
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        df = pd.read_excel(xlsx_path, sheet_name=spec["sheet"],
                           header=int(spec["header_row"]) - 1)
    rename, missing = {}, []
    for key, col in spec["columns"].items():
        if col in df.columns:
            rename[col] = key
        else:
            missing.append(col)
    if missing:
        raise ExtractError(
            f"商圏データシートに想定の列がありません: {missing}\n"
            f"  実際の列: {list(df.columns)[:12]}"
        )
    df = df.rename(columns=rename)[list(spec["columns"].keys())]
    return df.dropna(subset=["jicfs_mid_name"])


def extract_many(
    xlsx_paths: list[str | Path], map_path: str | Path, out_csv: str | Path
) -> pd.DataFrame:
    """複数ブックをまとめて1つの既存店マスタCSVにする。"""
    rows = [extract_store(p, map_path) for p in xlsx_paths]
    df = pd.DataFrame(rows)
    out_csv = Path(out_csv)
    out_csv.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(out_csv, index=False, encoding="utf-8-sig")
    return df
